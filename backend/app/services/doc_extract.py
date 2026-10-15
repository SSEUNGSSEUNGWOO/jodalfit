"""입찰 첨부파일(hwp / hwpx / pdf / xlsx, zip 묶음) 텍스트 추출.

외부 프로그램 없이 순수 Python. pyhwp는 유지보수 중단 + 무출력 문제로 쓰지 않고
HWP 5.0 공개 스펙대로 BodyText 문단 레코드를 직접 읽는다.
"""

from __future__ import annotations

import io
import re
import struct
import zipfile
import zlib
import xml.etree.ElementTree as ET

import olefile
from pypdf import PdfReader

SUPPORTED_EXTS = ("hwp", "hwpx", "pdf", "xlsx")
ARCHIVE_EXTS = ("zip",)


class UnsupportedDocument(Exception):
    """배포용(편집 제한)·암호 hwp 등 읽을 수 없는 문서."""


def ext_of(name: str | None) -> str:
    if not name or "." not in name:
        return ""
    return name.rsplit(".", 1)[1].strip().lower()


def sniff_format(data: bytes) -> str | None:
    """매직바이트로 실제 포맷 판별 — 나라장터 첨부는 확장자와 실제 포맷이 어긋나는 경우가 잦다
    (.hwpx인데 OLE hwp, .hwp인데 zip hwpx 등)."""
    head = data[:8]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "hwp"
    if head.startswith(b"PK\x03\x04"):
        return "hwpx"
    if head.lstrip().startswith(b"<?xml") and b"HWPML" in data[:2048]:
        return "hml"  # 한글 XML(HWPML) — .hwp/.hwpx 확장자로 올라오는 경우 있음
    return None


def extract_text(data: bytes, ext: str) -> str:
    # xlsx 는 zip 매직이 hwpx 와 겹쳐 sniff 로 가릴 수 없다 → 확장자 우선
    if ext == "xlsx":
        return extract_xlsx(data)
    fmt = sniff_format(data) or ext
    if fmt == "hwp":
        return extract_hwp(data)
    if fmt == "hwpx":
        return extract_hwpx(data)
    if fmt == "hml":
        return extract_hml(data)
    if fmt == "pdf":
        return extract_pdf(data)
    raise ValueError(f"unsupported format: ext={ext}")


# ----- HWPML (.hml, 단일 XML) -----

def extract_hml(data: bytes) -> str:
    root = ET.fromstring(data)
    out: list[str] = []
    for p in root.iter("P"):
        s = "".join(ch.text or "" for ch in p.iter("CHAR")).strip()
        if s:
            out.append(s)
    return _clean("\n".join(out))


# ----- PDF -----

def extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for pg in reader.pages:
        try:
            pages.append(pg.extract_text() or "")
        except Exception:  # 일부 페이지의 폰트 구조 오류(KeyError '/DescendantFonts' 등)는 그 페이지만 포기
            continue
    return _clean("\n".join(pages))


# ----- XLSX (규격서·내역서) -----

def extract_xlsx(data: bytes, max_rows_per_sheet: int = 2000) -> str:
    """시트별 텍스트. 셀은 탭, 행은 줄바꿈, 빈 셀·빈 행 생략. 수식은 값으로."""
    import openpyxl  # lazy — hwp/pdf 만 쓰는 경로의 import 비용 회피

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        lines: list[str] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows_per_sheet:
                lines.append("... (이하 생략)")  # read_only 모드에선 ws.max_row 가 None 일 수 있음
                break
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append("\t".join(cells))
        if lines:
            parts.append(f"[시트: {ws.title}]\n" + "\n".join(lines))
    wb.close()
    return "\n\n".join(parts)


# ----- ZIP (첨부 묶음) -----

def iter_zip_members(data: bytes) -> list[tuple[str, bytes]]:
    """zip 안의 지원 형식 파일만 (이름, 바이트) 로. 20MB 초과 멤버·암호 zip 은 건너뜀."""
    out: list[tuple[str, bytes]] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir() or info.file_size > 20 * 1024 * 1024:
                continue
            name = _decode_zip_name(info)
            if ext_of(name) in SUPPORTED_EXTS:
                try:
                    out.append((name, zf.read(info)))
                except RuntimeError:  # 암호
                    continue
    return out


def _decode_zip_name(info: zipfile.ZipInfo) -> str:
    """한글 파일명: UTF-8 플래그 없으면 cp949 로 복원 (나라장터 zip 대부분)."""
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode("cp949")
    except Exception:
        return info.filename


# ----- HWPX (OWPML, zip + xml) -----

_HP = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"


def extract_hwpx(data: bytes) -> str:
    out: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = sorted(
            n for n in z.namelist() if re.fullmatch(r"Contents/section\d+\.xml", n)
        )
        for n in names:
            root = ET.fromstring(z.read(n))
            # 표 셀 안 문단은 바깥 문단에 중첩됨 → 각 hp:t를 가장 가까운 hp:p에만 귀속시켜 중복 방지
            parent = {c: p for p in root.iter() for c in p}
            para_text: dict[ET.Element, list[str]] = {}
            order: list[ET.Element] = []
            for t in root.iter(f"{_HP}t"):
                p = parent.get(t)
                while p is not None and p.tag != f"{_HP}p":
                    p = parent.get(p)
                if p is None:
                    continue
                if p not in para_text:
                    para_text[p] = []
                    order.append(p)
                para_text[p].append(t.text or "")
            for p in order:
                s = "".join(para_text[p]).strip()
                if s:
                    out.append(s)
    return _clean("\n".join(out))


# ----- HWP 5.0 (OLE compound) -----

_HWPTAG_PARA_TEXT = 67  # HWPTAG_BEGIN(16) + 51
# 문단 텍스트 제어문자(코드 < 32) 중 1 wchar만 차지하는 문자형: 0, 줄바꿈(10), 문단끝(13), 24~31.
# 나머지(확장/인라인 컨트롤: 필드·표·탭 등)는 ch + 12바이트 + ch = 8 wchar 블록.
_CTRL_1WCHAR = {0, 10, 13, 24, 25, 26, 27, 28, 29, 30, 31}


def extract_hwp(data: bytes) -> str:
    ole = olefile.OleFileIO(io.BytesIO(data))
    fh = ole.openstream("FileHeader").read()
    if not fh.startswith(b"HWP Document File"):
        raise UnsupportedDocument("not an HWP 5.0 file")
    flags = struct.unpack("<I", fh[36:40])[0]
    if flags & 0x02:
        raise UnsupportedDocument("password protected")
    if flags & 0x04:
        raise UnsupportedDocument("distribution (배포용) document")
    compressed = bool(flags & 0x01)

    sections = sorted(
        (e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText"),
        key=lambda e: int(re.sub(r"\D", "", e[1]) or 0),
    )
    out: list[str] = []
    for entry in sections:
        body = ole.openstream(entry).read()
        if compressed:
            body = zlib.decompress(body, -15)
        for tag, payload in _iter_records(body):
            if tag != _HWPTAG_PARA_TEXT:
                continue
            s = _para_text(payload)
            if s:
                out.append(s)
    return _clean("\n".join(out))


def _iter_records(buf: bytes):
    i, n = 0, len(buf)
    while i + 4 <= n:
        hdr = struct.unpack_from("<I", buf, i)[0]
        i += 4
        tag = hdr & 0x3FF
        size = (hdr >> 20) & 0xFFF
        if size == 0xFFF:
            size = struct.unpack_from("<I", buf, i)[0]
            i += 4
        yield tag, buf[i : i + size]
        i += size


def _para_text(payload: bytes) -> str:
    s = payload.decode("utf-16le", "ignore")
    buf: list[str] = []
    j = 0
    while j < len(s):
        c = ord(s[j])
        if c < 32:
            if c in _CTRL_1WCHAR:
                if c in (10, 13):
                    buf.append("\n")
                elif c in (30, 31):
                    buf.append(" ")
                j += 1
            else:
                if c == 9:
                    buf.append("\t")
                j += 8
            continue
        buf.append(s[j])
        j += 1
    return "".join(buf).strip()


# ----- 공통 -----

def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[-]", "•", text)  # Wingdings/Symbol 글머리표(PUA) → 불릿
    lines = [re.sub(r"[ \t]+", " ", l).strip() for l in text.splitlines()]
    return "\n".join(l for l in lines if l)


if __name__ == "__main__":
    import openpyxl

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "sheet1"
    ws1.append(["a", "b"])
    ws1.append([None, None])  # 빈 행
    ws1.append(["c", None])  # 빈 셀
    wb.create_sheet("sheet2").append(["x"])
    buf = io.BytesIO()
    wb.save(buf)
    out = extract_xlsx(buf.getvalue())
    assert out.count("[시트:") == 2, out
    for line in out.splitlines():
        if not line.strip() or line.startswith("[시트:"):
            continue
        assert all(seg.strip() for seg in line.split("\t")), f"empty tab segment: {line!r}"

    zbuf = io.BytesIO()
    # zipfile 은 비-ascii 파일명을 항상 UTF-8 플래그로 쓴다 — 나라장터 zip 처럼 플래그 없이
    # cp437 바이트 그대로(=cp949로 재해석될 한글) 쓰려면 인코딩 단계를 잠시 바꿔치기한다.
    _orig_encode = zipfile.ZipInfo._encodeFilenameFlags
    zipfile.ZipInfo._encodeFilenameFlags = lambda self: (self.filename.encode("cp437"), self.flag_bits)
    try:
        with zipfile.ZipFile(zbuf, "w") as zf:
            zi = zipfile.ZipInfo("x")
            zi.filename = "규격서.xlsx".encode("cp949").decode("cp437")
            zf.writestr(zi, b"dummy")
            zf.writestr("note.txt", b"text file")
    finally:
        zipfile.ZipInfo._encodeFilenameFlags = _orig_encode
    members = iter_zip_members(zbuf.getvalue())
    names = [n for n, _ in members]
    assert "규격서.xlsx" in names, names
    assert not any(n.endswith(".txt") for n in names), names

    print("doc_extract self-check ok")
